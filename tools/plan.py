#!/usr/bin/env python3
"""Content plan operating layer for the CheckIA blog (messages are in French).

The source of truth is one brief per piece in content-plan/briefs/<id>-<slug>.md.
A brief has a small frontmatter (parsed here without PyYAML) that pre-answers the
AGENTS.md intake questions, plus a body with editorial notes and an append-only
« Journal ». Everything else (plan.csv for Google Sheets, plan-editorial.xlsx) is
generated from the briefs.

Usage:
  python3 tools/plan.py new <slug> --type texte --serie futur-de-l-audit --titre "…" [--mois 2026-11]
  python3 tools/plan.py build [--no-xlsx]
  python3 tools/plan.py list [--statut …] [--mois 2026-10] [--serie …] [--type …] [--responsable …]
  python3 tools/plan.py show P014
  python3 tools/plan.py next [--all]
  python3 tools/plan.py set P014 statut=redaction [cle=valeur …] [--note "…"] [--humain "Chloe"]
  python3 tools/plan.py check
  python3 tools/plan.py sync-sheet <published-csv-url | --file decisions.csv> [--dry-run]

Frontmatter subset: `key: value` (scalar), `key:` followed by `- item` lines (list),
`key: |` followed by lines indented by two spaces (block). No nesting, no quoting.
Human-only status transitions (brief-valide, pret-a-publier, archive) require --humain.
"""

import argparse
import csv
import html
import io
import re
import sys
import unicodedata
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SITE = "https://checkia.fr"
PLAN = ROOT / "content-plan"
BRIEFS = PLAN / "briefs"
CSV_PATH = PLAN / "plan.csv"
XLSX_PATH = PLAN / "plan-editorial.xlsx"
VERITE = PLAN / "produit-verite.md"
SUIVI = ROOT / "blog" / "suivi-ia.md"

SERIES = {
    "nouveautes-produit": "Nouveautés produit",
    "futur-de-l-audit": "Le futur de l'audit",
    "vie-de-l-entreprise": "Vie de l'entreprise",
    "temoignages-clients": "Témoignages clients",
    "site": "Pages du site",
}
TYPES = ["texte", "video", "temoignage", "refresh", "linkedin", "annuaire", "autre"]
ARTICLE_TYPES = {"texte", "video", "temoignage"}
TEMPLATE_SLUGS = {"modele-texte", "modele-temoignage", "ia-commissariat-aux-comptes"}
PRIORITES = ["haute", "normale", "basse"]

STATUTS = ["idee", "brief-valide", "en-attente", "redaction", "relecture",
           "pret-a-publier", "publie", "a-rafraichir", "archive"]
LABELS = {
    "idee": "Idée", "brief-valide": "Brief validé", "en-attente": "En attente (condition)",
    "redaction": "Rédaction", "relecture": "Relecture humaine",
    "pret-a-publier": "Prêt à publier", "publie": "Publié",
    "a-rafraichir": "À rafraîchir", "archive": "Archivé",
}
HUMAN_ONLY = {"brief-valide", "pret-a-publier", "archive"}
AGENT_ALLOWED = {
    ("idee", "idee"), ("brief-valide", "redaction"), ("brief-valide", "en-attente"),
    ("en-attente", "brief-valide"), ("redaction", "relecture"), ("relecture", "relecture"),
    ("pret-a-publier", "publie"), ("publie", "a-rafraichir"), ("a-rafraichir", "redaction"),
    ("publie", "publie"),
}
LIST_KEYS = {"requetes-secondaires", "questions", "faits", "chapitres", "liens-internes",
             "fonctionnalites", "conditions", "diffusion"}
ID_RE = re.compile(r"^P\d{3}$")
SLUG_RE = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
MOIS_RE = re.compile(r"^\d{4}-\d{2}$")

# (frontmatter key, CSV header). `id` first: it is the join key in Google Sheets.
CSV_COLUMNS = [
    ("id", "id"), ("statut", "Statut"), ("type", "Type"), ("serie", "Série"),
    ("rubrique", "Rubrique"), ("pilier", "Pilier"), ("titre", "Titre de travail"),
    ("requete", "Requête cible"), ("requetes-secondaires", "Requêtes secondaires"),
    ("auteur", "Auteur"), ("interviewe", "Interviewé"), ("mois", "Mois"),
    ("date-prevue", "Date prévue"), ("date-publiee", "Date publiée"), ("url", "URL"),
    ("priorite", "Priorité"), ("responsable", "Responsable"), ("image-og", "Image OG"),
    ("video-id", "Vidéo (ID)"), ("conditions", "Conditions"), ("rafraichit", "Rafraîchit"),
    ("piece-liee", "Pièce liée"), ("fonctionnalites", "Fonctionnalités"),
    ("diffusion", "Diffusion"), ("angle", "Angle"), ("fichier", "Fichier du brief"),
    ("indexe", "Indexé (site)"),
]
DECISION_COLUMNS = ["id", "Statut décidé", "Commentaire", "Décidé par", "Date"]

errors, warnings = [], []


def err(msg):
    errors.append("[ERREUR] " + msg)


def warn(msg):
    warnings.append("[attention] " + msg)


def fold(s):
    """Lowercase, strip accents and punctuation (same rule as check-seo.py)."""
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def today():
    return date.today().isoformat()


# --------------------------------------------------------------------------- frontmatter

FRONT_RE = re.compile(r"\A---\n(.*?)\n---\n?(.*)\Z", re.S)


def parse_front(text):
    """Parse the frontmatter subset. Returns (dict, body).

    Values are str (scalar or block with newlines) or list[str]."""
    m = FRONT_RE.match(text)
    if not m:
        raise ValueError("frontmatter absent (--- … ---)")
    data, body = {}, m.group(2)
    key, mode = None, None
    for raw in m.group(1).splitlines():
        if mode == "block" and (raw.startswith("  ") or raw.strip() == ""):
            data[key] += raw[2:] + "\n"
            continue
        s = raw.strip()
        if mode == "list" and s.startswith("- "):
            data[key].append(s[2:].strip())
            continue
        if not s or s.startswith("#"):
            mode = None
            continue
        if ":" not in s:
            raise ValueError("ligne invalide dans le frontmatter : %r" % raw)
        key, _, val = s.partition(":")
        key, val = key.strip().lower(), val.strip()
        if val == "":
            data[key], mode = [], "list"
        elif val == "|":
            data[key], mode = "", "block"
        else:
            data[key], mode = ("" if val in ("~", '""') else val), None
    clean = {}
    for k, v in data.items():
        if isinstance(v, str):
            v = v.rstrip("\n")
        if isinstance(v, list) and k not in LIST_KEYS and not v:
            v = ""  # `key:` with nothing after it is an empty scalar for non-list keys
        clean[k] = v
    return clean, body


def dump_front(data, order):
    lines = []
    keys = [k for k in order if k in data] + [k for k in data if k not in order]
    for k in keys:
        v = data[k]
        if isinstance(v, list):
            lines.append(k + ":")
            lines += ["  - " + i for i in v]
        elif "\n" in v:
            lines.append(k + ": |")
            lines += ["  " + l for l in v.split("\n")]
        else:
            lines.append(k + ": " + v if v else k + ":")
    return "---\n" + "\n".join(lines) + "\n---\n"


class Brief:
    def __init__(self, path):
        self.path = path
        text = path.read_text(encoding="utf-8")
        self.data, self.body = parse_front(text)
        self.order = list(self.data)

    @property
    def id(self):
        return self.data.get("id", "")

    def get(self, k, default=""):
        v = self.data.get(k, default)
        if isinstance(v, list):
            return v
        return v if v is not None else default

    def lst(self, k):
        v = self.data.get(k, [])
        if isinstance(v, str):
            return [x.strip() for x in v.split(" ; ") if x.strip()] if v else []
        return v

    @property
    def statut(self):
        return self.get("statut", "idee")

    @property
    def type(self):
        return self.get("type")

    def is_article(self):
        return self.type in ARTICLE_TYPES and self.get("serie") and self.get("serie") != "site"

    def article_dir(self):
        if self.get("serie") == "site" and self.get("slug"):
            return ROOT / self.get("slug")
        if not self.is_article() or not self.get("slug"):
            return None
        return ROOT / "blog" / self.get("serie") / self.get("slug")

    def article_html(self):
        d = self.article_dir()
        if d is None:
            return None
        p = d / "index.html"
        return p.read_text(encoding="utf-8") if p.exists() else None

    def is_indexed(self):
        t = self.article_html()
        if t is None:
            return None
        t = re.sub(r"<!--.*?-->", "", t, flags=re.S)
        return not re.search(r'<meta name="robots" content="[^"]*noindex', t)

    def canonical(self):
        d = self.article_dir()
        if d is None:
            return ""
        return SITE + "/" + str(d.relative_to(ROOT)).replace("\\", "/") + "/"

    def conditions_ok(self):
        return all(fold(c).endswith(" ok") for c in self.lst("conditions"))

    def append_journal(self, line):
        entry = "- " + line
        if "## Journal" in self.body:
            self.body = self.body.rstrip("\n") + "\n" + entry + "\n"
        else:
            self.body = self.body.rstrip("\n") + "\n\n## Journal\n\n" + entry + "\n"

    def save(self, note=None, who="agent"):
        if note:
            self.append_journal("%s — %s — %s" % (today(), who, note))
        self.path.write_text(dump_front(self.data, self.order) + self.body, encoding="utf-8")


def load_briefs():
    BRIEFS.mkdir(parents=True, exist_ok=True)
    out = []
    for p in sorted(BRIEFS.glob("*.md")):
        try:
            out.append(Brief(p))
        except ValueError as e:
            err("%s : %s" % (p.relative_to(ROOT), e))
    return out


def find_brief(bid):
    bid = bid.strip().upper()
    for b in load_briefs():
        if b.id == bid:
            return b
    sys.exit("Brief %s introuvable dans content-plan/briefs/." % bid)


def normalize_statut(value):
    v = fold(value)
    for slug, label in LABELS.items():
        if v in (fold(slug), fold(label)):
            return slug
    sys.exit("Statut inconnu : « %s ». Choix : %s" % (value, ", ".join(LABELS.values())))


def pilier(b):
    """Editorial pillar, derived from série / rubrique / type for the Piliers tab."""
    if b.get("rubrique") == "paroles-de-cac" or b.type == "temoignage":
        return "P4 Paroles de CAC"
    if b.type in ("linkedin", "annuaire", "autre"):
        return "Hors site"
    p = b.get("pilier")
    if p:
        return p
    return {"futur-de-l-audit": "P1 IA et CAC", "nouveautes-produit": "P5 Produit",
            "vie-de-l-entreprise": "P6 Entreprise", "site": "P6 Entreprise"}.get(b.get("serie"), "")


# --------------------------------------------------------------------------- produit-verite

def load_verite():
    """Return {section: [feature ids]} from produit-verite.md (`- [id] …` lines)."""
    out = {}
    if not VERITE.exists():
        return out
    section = None
    for line in VERITE.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            section = line[3:].strip()
            out.setdefault(section, [])
        elif section and line.startswith("- ["):
            m = re.match(r"- \[([a-z0-9-]+)\]", line)
            if m:
                out[section].append(m.group(1))
    return out


def verite_sets():
    v = load_verite()
    prod = {i for s, ids in v.items() if fold(s).startswith("en production") for i in ids}
    road = {i for s, ids in v.items() if fold(s).startswith("feuille de route") for i in ids}
    return prod, road


# --------------------------------------------------------------------------- commands

def cmd_new(args):
    briefs = load_briefs()
    nums = [int(b.id[1:]) for b in briefs if ID_RE.match(b.id)]
    bid = "P%03d" % (max(nums) + 1 if nums else 1)
    if not SLUG_RE.match(args.slug):
        sys.exit("Slug invalide : minuscules, chiffres et tirets, sans accents.")
    if args.type not in TYPES:
        sys.exit("Type inconnu : %s. Choix : %s" % (args.type, ", ".join(TYPES)))
    if args.type in ARTICLE_TYPES and args.serie not in SERIES:
        sys.exit("Série requise pour un article : %s" % ", ".join(SERIES))
    path = BRIEFS / ("%s-%s.md" % (bid, args.slug))
    front = {
        "id": bid, "slug": args.slug, "type": args.type, "serie": args.serie or "",
        "rubrique": "", "pilier": "", "titre": args.titre, "angle": "", "description": "",
        "requete": "aucune", "requetes-secondaires": [], "questions": [], "faits": [],
        "auteur": args.auteur or "L'équipe CheckIA", "auteur-titre": "", "auteur-linkedin": "",
        "interviewe": "", "interviewe-titre": "", "interviewe-linkedin": "",
        "image-og": "generee", "cta": "appel-30-min", "video-id": "", "video-duree": "",
        "chapitres": [], "liens-internes": [], "piece-liee": "", "rafraichit": "",
        "fonctionnalites": [], "conditions": [], "statut": "idee", "priorite": "normale",
        "mois": args.mois or "", "date-prevue": "", "date-publiee": "", "url": "",
        "diffusion": [], "responsable": args.responsable or "Chloe",
    }
    body = ("\n## Angle éditorial\n\n(à rédiger)\n\n## Plan pressenti\n\n1. \n\n## Sources\n\n- \n\n"
            "## Journal\n\n- %s — %s — statut : idee (création du brief)\n" % (today(), args.responsable or "agent"))
    path.write_text(dump_front(front, list(front)) + body, encoding="utf-8")
    print("Créé : %s (statut : Idée)" % path.relative_to(ROOT))


def brief_row(b):
    row = {}
    for key, header in CSV_COLUMNS:
        if key == "statut":
            v = LABELS.get(b.statut, b.statut)
        elif key == "pilier":
            v = pilier(b)
        elif key == "fichier":
            v = str(b.path.relative_to(ROOT))
        elif key == "indexe":
            ix = b.is_indexed()
            v = "—" if ix is None else ("oui" if ix else "non")
        elif key == "serie":
            v = SERIES.get(b.get("serie"), b.get("serie"))
        else:
            v = b.get(key)
        if isinstance(v, list):
            v = " ; ".join(v)
        row[header] = (v or "").replace("\n", " / ")
    return row


def sort_key(b):
    return (b.get("mois") or "9999-99", b.get("date-prevue") or "9999-99-99", b.id)


def build_csv(briefs, path=CSV_PATH):
    headers = [h for _, h in CSV_COLUMNS]
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=headers, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        for b in sorted(briefs, key=sort_key):
            w.writerow(brief_row(b))
    print("Écrit : %s (%d lignes)" % (path.relative_to(ROOT), len(briefs)))


def parse_suivi():
    """Rows of the markdown table in blog/suivi-ia.md (list of lists)."""
    rows = []
    if not SUIVI.exists():
        return rows
    for line in SUIVI.read_text(encoding="utf-8").splitlines():
        if line.startswith("|") and not re.match(r"^\|\s*-", line):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            rows.append(cells)
    return rows


def due_checks(briefs):
    """(brief, echeance, due_date) for publie pieces missing a +7 / +30 day citation check."""
    logged = parse_suivi()[1:]
    out = []
    for b in briefs:
        if b.statut != "publie" or not b.get("date-publiee"):
            continue
        try:
            pub = datetime.strptime(b.get("date-publiee"), "%Y-%m-%d").date()
        except ValueError:
            continue
        url = b.get("url") or b.canonical()
        for label, days in (("J+7", 7), ("J+30", 30)):
            due = pub + timedelta(days=days)
            if due > date.today():
                continue
            has = any(len(r) > 1 and url and url in r[1]
                      and r[0][:10] >= (due - timedelta(days=3)).isoformat() for r in logged)
            if not has:
                out.append((b, label, due.isoformat()))
    return out


def build_xlsx(briefs, path=XLSX_PATH):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        warn("openpyxl absent : plan-editorial.xlsx non généré (pip3 install openpyxl)")
        return
    fills = {
        "Idée": "EEEEEE", "Brief validé": "DDEBF7", "En attente (condition)": "FCE4D6",
        "Rédaction": "FFF2CC", "Relecture humaine": "FFE699", "Prêt à publier": "C6E0B4",
        "Publié": "A9D08E", "À rafraîchir": "F8CBAD", "Archivé": "D9D9D9",
    }
    bold = Font(bold=True)

    def sheet(ws, headers, rows, widths=None, status_col=None):
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
            if status_col is not None and r[status_col] in fills:
                ws.cell(row=ws.max_row, column=status_col + 1).fill = PatternFill(
                    "solid", fgColor=fills[r[status_col]])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, 1):
            w = (widths or {}).get(h, 18)
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    headers = [h for _, h in CSV_COLUMNS]
    ordered = sorted(briefs, key=sort_key)
    rows, last = [], None
    for b in ordered:
        r = brief_row(b)
        if last is not None and r["Mois"] != last:
            rows.append([""] * len(headers))
        last = r["Mois"]
        rows.append([r[h] for h in headers])
    sheet(wb.active, headers, rows,
          {"Titre de travail": 48, "Angle": 60, "Requête cible": 28, "Conditions": 30,
           "Requêtes secondaires": 30, "URL": 40, "Fichier du brief": 40}, status_col=1)
    wb.active.title = "Calendrier"

    # Piliers : pilier × mois
    months = sorted({b.get("mois") for b in briefs if b.get("mois")})
    pils = sorted({pilier(b) for b in briefs})
    ws = wb.create_sheet("Piliers")
    prow = []
    for p in pils:
        counts = [sum(1 for b in briefs if pilier(b) == p and b.get("mois") == m) for m in months]
        prow.append([p] + counts + [sum(counts)])
    sheet(ws, ["Pilier"] + months + ["Total"], prow, {"Pilier": 28})
    ws.append([])
    ws.append(["Type"] + months + ["Total"])
    for t in TYPES:
        counts = [sum(1 for b in briefs if b.type == t and b.get("mois") == m) for m in months]
        ws.append([t] + counts + [sum(counts)])
    ws.append([])
    ws.append(["Pilier", "id", "Titre", "Statut"])
    for p in pils:
        for b in [x for x in ordered if pilier(x) == p]:
            ws.append([p, b.id, b.get("titre"), LABELS.get(b.statut, b.statut)])

    # Requêtes
    ws = wb.create_sheet("Requêtes")
    qrows, seen = [], {}
    for b in ordered:
        if b.statut == "archive":
            continue
        q = b.get("requete")
        if q and fold(q) != "aucune":
            seen.setdefault(fold(q), []).append(b.id)
    for b in ordered:
        if b.statut == "archive":
            continue
        entries = []
        q = b.get("requete")
        if q and fold(q) != "aucune":
            entries.append((q, "cible"))
        entries += [(s, "secondaire") for s in b.lst("requetes-secondaires")]
        for q, role in entries:
            in_title = "—"
            html_text = b.article_html()
            if html_text and role == "cible":
                m = re.search(r"<title>([^<]*)</title>", html_text)
                if m:
                    tw = fold(html.unescape(m.group(1))).split()
                    in_title = "oui" if all(w in tw for w in fold(q).split()) else "non"
            dup = "doublon" if role == "cible" and len(seen.get(fold(q), [])) > 1 else ""
            qrows.append([q, role, b.id, b.get("titre"), LABELS.get(b.statut, b.statut),
                          b.get("url"), dup, in_title])
    sheet(ws, ["Requête", "Rôle", "id", "Titre", "Statut", "URL", "Doublon", "Requête dans le <title>"],
          qrows, {"Requête": 40, "Titre": 48, "URL": 40})

    # Distribution
    ws = wb.create_sheet("Distribution")
    channels = ["linkedin", "x", "instagram", "indexnow", "search-console"]
    drows = []
    for b in ordered:
        if b.statut not in ("relecture", "pret-a-publier", "publie", "a-rafraichir"):
            continue
        done = {}
        for d in b.lst("diffusion"):
            ch, _, when = d.partition("—")
            done[fold(ch).replace(" ", "-")] = when.strip()
        drows.append([b.id, b.get("titre"), b.get("url") or b.canonical(), b.get("date-publiee")]
                     + [done.get(c, "") for c in channels])
    sheet(ws, ["id", "Titre", "URL", "Date publiée", "LinkedIn", "X", "Instagram", "IndexNow", "Search Console"],
          drows, {"Titre": 48, "URL": 40})

    # Suivi IA
    ws = wb.create_sheet("Suivi IA")
    suivi = parse_suivi()
    if suivi:
        sheet(ws, suivi[0], suivi[1:], {h: 30 for h in suivi[0]})
    else:
        sheet(ws, ["Date", "Article (URL)", "Requête posée", "Moteur", "Cité ?", "Page citée",
                   "Concurrent cité à la place", "Action décidée"], [])
    ws.append([])
    ws.append(["Vérifications dues", "id", "Titre", "Échéance", "Date due"])
    for b, label, due in due_checks(briefs):
        ws.append(["", b.id, b.get("titre"), label, due])

    # Backlog
    ws = wb.create_sheet("Backlog")
    brows = [[b.id, LABELS.get(b.statut, b.statut), b.get("titre"), pilier(b), b.get("priorite"),
              b.get("responsable"), " ; ".join(b.lst("conditions"))]
             for b in ordered if b.statut in ("idee", "en-attente")]
    sheet(ws, ["id", "Statut", "Titre", "Pilier", "Priorité", "Responsable", "Conditions"],
          brows, {"Titre": 48, "Conditions": 40}, status_col=1)

    # Légende
    ws = wb.create_sheet("Légende")
    who = {"idee": "tout le monde", "brief-valide": "humain", "en-attente": "agent ou humain",
           "redaction": "agent", "relecture": "agent", "pret-a-publier": "humain",
           "publie": "agent (après mise en ligne)", "a-rafraichir": "humain ou agent (motif au Journal)",
           "archive": "humain"}
    does = {"idee": "Enrichit le brief (requêtes, questions, sources). Rien dans blog/.",
            "brief-valide": "Lit le corpus, génère le squelette, commence la rédaction.",
            "en-attente": "Ne fait rien sur la pièce ; revérifie les conditions.",
            "redaction": "Rédige l'article en noindex, images, liens, build-llms, check-seo, social.md.",
            "relecture": "Applique les retours ; ne passe jamais robots en index.",
            "pret-a-publier": "Exécute la mise en ligne (index, sitemap, flux, llms.txt, commit).",
            "publie": "Vérifie les citations IA à J+7 et J+30 (blog/suivi-ia.md).",
            "a-rafraichir": "Prépare le rafraîchissement, soumet le diff en relecture.",
            "archive": "Rien."}
    sheet(ws, ["Statut", "Qui décide", "Ce que fait l'agent ensuite"],
          [[LABELS[s], who[s], does[s]] for s in STATUTS], {"Statut": 24, "Qui décide": 30, "Ce que fait l'agent ensuite": 80})
    ws.append([])
    ws.append(["Types", ", ".join(TYPES)])
    ws.append(["Séries", ", ".join("%s (%s)" % (k, v) for k, v in SERIES.items())])
    ws.append([])
    ws.append(["Mode d'emploi"])
    for line in [
        "1. L'onglet « Plan (repo) » est un miroir du dépôt (IMPORTDATA) : ne pas le modifier.",
        "2. Décider dans l'onglet « Décisions » : id, statut décidé, commentaire, prénom, date. Une ligne par décision, jamais d'édition des anciennes lignes.",
        "3. « Brief validé » lance la rédaction ; « Prêt à publier » vaut feu vert de publication.",
        "4. Pour forcer le rafraîchissement du miroir : incrémenter la cellule Version de l'onglet Réglages.",
        "5. Les détails d'un brief (faits, questions, sources) sont dans content-plan/briefs/ sur GitHub.",
    ]:
        ws.append([line])
    wb.save(path)
    print("Écrit : %s (7 onglets)" % path.relative_to(ROOT))


def cmd_build(args):
    briefs = load_briefs()
    if errors:
        print("\n".join(errors))
        sys.exit(1)
    build_csv(briefs)
    if not args.no_xlsx:
        build_xlsx(briefs)
    for w in warnings:
        print(w)


def cmd_list(args):
    briefs = load_briefs()
    rows = []
    for b in sorted(briefs, key=sort_key):
        if args.statut and b.statut != normalize_statut(args.statut):
            continue
        if args.mois and b.get("mois") != args.mois:
            continue
        if args.serie and b.get("serie") != args.serie:
            continue
        if args.type and b.type != args.type:
            continue
        if args.responsable and fold(b.get("responsable")) != fold(args.responsable):
            continue
        rows.append((b.id, LABELS.get(b.statut, b.statut), b.get("mois"), b.type,
                     b.get("serie"), b.get("titre")))
    for r in rows:
        print("%-5s %-24s %-8s %-11s %-22s %s" % r)
    print("\n%d pièce(s)." % len(rows))


def cmd_show(args):
    b = find_brief(args.id)
    labels = {k: h for k, h in CSV_COLUMNS}
    for k in b.order:
        v = b.data[k]
        label = labels.get(k, k)
        if isinstance(v, list):
            print("%s :" % label)
            for i in v:
                print("  - " + i)
        elif v:
            print("%s : %s" % (label, LABELS.get(v, v) if k == "statut" else v))
    print(b.body)


def journal_lines(b):
    m = re.search(r"## Journal\n(.*)\Z", b.body, re.S)
    return [l for l in (m.group(1) if m else "").splitlines() if l.startswith("- ")]


def cmd_next(args):
    briefs = load_briefs()
    items = []
    for b in sorted(briefs, key=sort_key):
        if b.statut == "pret-a-publier":
            items.append((0, b, "Publier (mise en ligne AGENTS.md « Going live »)",
                          "python3 tools/check-seo.py ; puis robots index, sitemap, feed, llms.txt, build-llms, "
                          "plan.py set %s statut=publie" % b.id))
    for b in sorted(briefs, key=sort_key):
        if b.statut == "relecture":
            lines = journal_lines(b)
            if lines and "agent" not in lines[-1]:
                items.append((1, b, "Appliquer les retours de relecture (dernière ligne du Journal)",
                              "python3 tools/plan.py show %s" % b.id))
    for b in sorted(briefs, key=lambda x: (x.get("date-prevue") or "9999", PRIORITES.index(x.get("priorite") or "normale"), x.id)):
        if b.statut == "brief-valide" and b.conditions_ok():
            cmd = ("python3 tools/plan.py set %s statut=redaction && python3 tools/new-article.py --brief %s"
                   % (b.id, b.id)) if b.is_article() else "python3 tools/plan.py set %s statut=redaction" % b.id
            items.append((2, b, "Rédiger", cmd))
    for b in sorted(briefs, key=sort_key):
        if b.statut == "a-rafraichir":
            items.append((3, b, "Préparer le rafraîchissement", "python3 tools/plan.py show %s" % b.id))
    for b, label, due in due_checks(briefs):
        items.append((4, b, "Vérifier les citations IA (%s, échéance %s)" % (label, due),
                      "Procédure « Citation check » (AGENTS.md) → blog/suivi-ia.md"))
    for b in sorted(briefs, key=sort_key):
        if b.statut == "brief-valide" and not b.conditions_ok():
            pending = [c for c in b.lst("conditions") if not fold(c).endswith(" ok")]
            items.append((5, b, "Brief validé mais bloqué par une condition : demander à l'humain — %s" % " ; ".join(pending),
                          "python3 tools/plan.py set %s conditions=\"… — ok\" (après confirmation humaine)" % b.id))
    if not items:
        print("Rien à faire : aucune pièce en attente d'action de l'agent.")
        return
    for _, b, why, cmd in (items if args.all else items[:1]):
        print("%s — %s\n  Action : %s\n  Commande : %s\n" % (b.id, b.get("titre"), why, cmd))


def cmd_set(args):
    b = find_brief(args.id)
    changes = {}
    for kv in args.assignments:
        if "=" not in kv:
            sys.exit("Affectation attendue : cle=valeur (reçu : %s)" % kv)
        k, v = kv.split("=", 1)
        changes[k.strip().lower()] = v.strip()
    note = args.note
    if "statut" in changes:
        new = normalize_statut(changes["statut"])
        old = b.statut
        if new in HUMAN_ONLY and not args.humain:
            sys.exit("Transition %s → %s réservée à un humain : ajouter --humain \"Prénom\" après "
                     "accord explicite dans la conversation." % (old, new))
        if not args.humain and (old, new) not in AGENT_ALLOWED:
            sys.exit("Transition %s → %s non autorisée pour l'agent." % (old, new))
        if new == "brief-valide" and old == "en-attente" and not b.conditions_ok():
            sys.exit("Conditions non remplies : %s" % " ; ".join(b.lst("conditions")))
        if new == "publie":
            if b.is_article() and b.is_indexed() is not True:
                sys.exit("L'article n'est pas indexé (robots) : terminer la mise en ligne avant statut=publie.")
            changes.setdefault("date-publiee", today())
            if b.is_article():
                changes.setdefault("url", b.canonical())
        changes["statut"] = new
        note = note or "statut : %s → %s" % (old, new)
    for k, v in changes.items():
        if k in LIST_KEYS:
            b.data[k] = [x.strip() for x in v.split(" ; ") if x.strip()]
        else:
            b.data[k] = v
        if k not in b.order:
            b.order.append(k)
    b.save(note=note or "mise à jour : " + ", ".join(changes), who=args.humain or "agent")
    print("Mis à jour : %s (%s)" % (b.id, ", ".join("%s=%s" % kv for kv in changes.items())))
    if args.build:
        build_csv(load_briefs())
        build_xlsx(load_briefs())


def cmd_check(args):
    briefs = load_briefs()
    prod, road = verite_sets()
    if not VERITE.exists():
        warn("content-plan/produit-verite.md absent")
    ids = {}
    slugs, queries = {}, {}
    rank = {s: i for i, s in enumerate(STATUTS)}
    for b in briefs:
        rel = str(b.path.relative_to(ROOT))
        bid = b.id
        if not ID_RE.match(bid):
            err("%s : id invalide « %s »" % (rel, bid))
            continue
        if not b.path.name.startswith(bid + "-"):
            err("%s : le nom du fichier ne commence pas par %s-" % (rel, bid))
        if bid in ids:
            err("%s : id %s en doublon avec %s" % (rel, bid, ids[bid]))
        ids[bid] = rel
    by_id = {b.id: b for b in briefs}
    for b in briefs:
        p = b.id
        st = b.statut
        if st not in STATUTS:
            err("%s : statut inconnu « %s »" % (p, st))
            continue
        if b.type not in TYPES:
            err("%s : type inconnu « %s »" % (p, b.type))
        if b.get("priorite") and b.get("priorite") not in PRIORITES:
            err("%s : priorité inconnue « %s »" % (p, b.get("priorite")))
        serie = b.get("serie")
        if b.type in ARTICLE_TYPES and serie not in SERIES:
            err("%s : série requise pour un article (%s)" % (p, ", ".join(SERIES)))
        if b.type in ("linkedin", "annuaire") and serie:
            warn("%s : série renseignée pour une pièce hors site" % p)
        slug = b.get("slug")
        if slug and not SLUG_RE.match(slug):
            err("%s : slug invalide « %s »" % (p, slug))
        if slug and st != "archive" and b.type in ARTICLE_TYPES:
            key = (serie, slug)
            if key in slugs:
                err("%s : slug « %s » en doublon avec %s" % (p, slug, slugs[key]))
            slugs[key] = p
        for k in ("date-prevue", "date-publiee"):
            if b.get(k) and not DATE_RE.match(b.get(k)):
                err("%s : %s non ISO (AAAA-MM-JJ) : %s" % (p, k, b.get(k)))
        if b.get("mois") and not MOIS_RE.match(b.get("mois")):
            err("%s : mois attendu au format AAAA-MM : %s" % (p, b.get("mois")))
        if b.get("mois") and b.get("date-prevue") and b.get("date-prevue")[:7] != b.get("mois"):
            warn("%s : mois %s ≠ date prévue %s" % (p, b.get("mois"), b.get("date-prevue")))
        q = b.get("requete")
        if q and fold(q) != "aucune" and st != "archive":
            if fold(q) in queries:
                err("%s : requête cible « %s » déjà visée par %s" % (p, q, queries[fold(q)]))
            queries[fold(q)] = p
        if rank[st] >= rank["brief-valide"] and st != "en-attente":
            for k in ("titre", "angle", "auteur", "image-og"):
                if not b.get(k):
                    err("%s : champ « %s » requis à partir de Brief validé" % (p, k))
            if b.type in ARTICLE_TYPES or b.type == "refresh":
                faits = b.lst("faits")
                if len(faits) < 3:
                    err("%s : au moins 3 faits sourcés requis (%d)" % (p, len(faits)))
                for f in faits:
                    if "source" not in fold(f):
                        err("%s : fait sans source : « %s »" % (p, f[:60]))
                if b.type in ARTICLE_TYPES and len(b.lst("questions")) < 2:
                    err("%s : au moins 2 questions (h2 / FAQ) requises" % p)
        if b.type in ("video", "temoignage") and rank[st] >= rank["relecture"] and st != "a-rafraichir":
            if not b.get("video-id") or not b.get("video-duree"):
                err("%s : video-id et video-duree requis à partir de Relecture" % p)
        for k in ("rafraichit", "piece-liee"):
            if b.get(k) and b.get(k) not in by_id:
                err("%s : %s pointe vers un brief inconnu (%s)" % (p, k, b.get(k)))
        for l in b.lst("liens-internes"):
            if ID_RE.match(l) and l not in by_id:
                err("%s : lien interne vers un brief inconnu (%s)" % (p, l))
        if b.type == "refresh":
            t = b.get("rafraichit")
            if not t:
                err("%s : type refresh sans « rafraichit »" % p)
            elif by_id[t].statut not in ("publie", "a-rafraichir"):
                warn("%s : la pièce rafraîchie %s n'est pas publiée" % (p, t))
        if b.get("rubrique") == "paroles-de-cac":
            if serie != "futur-de-l-audit":
                err("%s : la rubrique paroles-de-cac vit dans futur-de-l-audit" % p)
            if b.type not in ("video", "texte"):
                err("%s : paroles-de-cac : type video ou texte" % p)
            if rank[st] >= rank["brief-valide"] and not b.get("interviewe"):
                err("%s : interviewé requis pour paroles-de-cac" % p)
        feats = b.lst("fonctionnalites")
        if serie == "nouveautes-produit" and rank[st] >= rank["redaction"] and not feats:
            warn("%s : article produit sans « fonctionnalites » (ids de produit-verite.md)" % p)
        for fid in feats:
            if fid in prod:
                continue
            if fid in road:
                (err if rank[st] >= rank["redaction"] and st != "en-attente" else warn)(
                    "%s : fonctionnalité « %s » en feuille de route seulement" % (p, fid))
            else:
                err("%s : fonctionnalité « %s » absente de produit-verite.md" % (p, fid))
        if st == "publie":
            if not b.get("date-publiee") or not b.get("url"):
                err("%s : Publié sans date-publiee / url" % p)
        # Cross-check with the site
        d = b.article_dir()
        if d is not None:
            exists = (d / "index.html").exists()
            ix = b.is_indexed()
            if st == "publie" and not exists:
                err("%s : Publié mais %s/index.html absent" % (p, d.relative_to(ROOT)))
            elif st == "publie" and ix is False:
                err("%s : Publié mais la page est en noindex" % p)
            elif st in ("redaction", "relecture", "pret-a-publier") and exists and ix:
                err("%s : statut %s mais la page est déjà indexée" % (p, st))
            elif st in ("idee", "brief-valide") and exists and b.type in ARTICLE_TYPES and serie != "site":
                warn("%s : la page existe déjà alors que le brief est en %s" % (p, LABELS[st]))
            if exists and serie != "site":
                text = (d / "index.html").read_text(encoding="utf-8")
                meta = re.search(r"<!--\s*checkia-meta\s*(.*?)-->", text, re.S)
                md = {}
                if meta:
                    for line in meta.group(1).splitlines():
                        if ":" in line:
                            k, v = line.split(":", 1)
                            md[k.strip().lower()] = re.sub(r"\s*\(.*\)\s*$", "", v).strip()
                if md.get("brief") != p:
                    warn("%s : bloc checkia-meta sans « brief: %s »" % (p, p))
                if md.get("format") and md.get("format") != b.type:
                    err("%s : format « %s » ≠ type du brief « %s »" % (p, md["format"], b.type))
                mq = md.get("query", "none")
                if fold(mq if mq.lower() != "none" else "aucune") != fold(q or "aucune"):
                    err("%s : query du bloc meta « %s » ≠ requête du brief « %s »" % (p, mq, q))
                if md.get("author") and fold(md["author"]) != fold(b.get("auteur")):
                    err("%s : author « %s » ≠ auteur du brief « %s »" % (p, md["author"], b.get("auteur")))
                if q and fold(q) != "aucune":
                    m = re.search(r"<title>([^<]*)</title>", re.sub(r"<!--.*?-->", "", text, flags=re.S))
                    if m:
                        tw = fold(html.unescape(m.group(1))).split()
                        missing = [w for w in fold(q).split() if w not in tw]
                        if missing:
                            (err if rank[st] >= rank["relecture"] else warn)(
                                "%s : mots de la requête absents du <title> : %s" % (p, ", ".join(missing)))
    # Reverse: every blog article has a brief
    briefed = {(b.get("serie"), b.get("slug")) for b in briefs if b.is_article()}
    for page in sorted(ROOT.glob("blog/*/*/index.html")):
        serie, slug = page.parts[-3], page.parts[-2]
        if slug in TEMPLATE_SLUGS:
            continue
        text = re.sub(r"<!--.*?-->", "", page.read_text(encoding="utf-8"), flags=re.S)
        indexed = not re.search(r'<meta name="robots" content="[^"]*noindex', text)
        if (serie, slug) not in briefed:
            (err if indexed else warn)("blog/%s/%s : article sans brief (plan.py new %s …)" % (serie, slug, slug))
        elif indexed:
            b = next(x for x in briefs if x.get("serie") == serie and x.get("slug") == slug)
            if b.statut != "publie":
                err("blog/%s/%s : page indexée mais brief %s en statut %s" % (serie, slug, b.id, LABELS[b.statut]))
    if CSV_PATH.exists() and briefs:
        newest = max(b.path.stat().st_mtime for b in briefs)
        if CSV_PATH.stat().st_mtime < newest:
            warn("plan.csv plus ancien que le dernier brief : lancer python3 tools/plan.py build")
    for w in warnings:
        print(w)
    for e in errors:
        print(e)
    print("\n%d brief(s) vérifié(s) — %d erreur(s), %d avertissement(s)." % (len(briefs), len(errors), len(warnings)))
    sys.exit(1 if errors else 0)


def fetch_decisions(source):
    if source.startswith("http://") or source.startswith("https://"):
        with urllib.request.urlopen(source, timeout=30) as r:
            raw = r.read()
    else:
        raw = Path(source).read_bytes()
    text = raw.decode("utf-8-sig")
    sample = text[:2048]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    try:
        header = next(reader)
    except StopIteration:
        return []
    # Tolerant header mapping: accents/case ignored, « Statut » alone accepted for
    # « Statut décidé », stray cells (e.g. a split « décidé ») ignored.
    aliases = {"id": "id", "statut decide": "Statut décidé", "statut": "Statut décidé",
               "commentaire": "Commentaire", "decide par": "Décidé par", "date": "Date"}
    index = {}
    for i, h in enumerate(header):
        name = aliases.get(fold(h))
        if name and name not in index:
            index[name] = i
    missing = [c for c in DECISION_COLUMNS if c not in index]
    if missing:
        sys.exit("Colonnes manquantes dans la feuille Décisions : %s (en-têtes lus : %s)"
                 % (", ".join(missing), ", ".join(header)))
    rows = []
    for cells in reader:
        if not any(c.strip() for c in cells):
            continue
        rows.append({c: (cells[i] if i < len(cells) else "") for c, i in index.items()})
    return rows


def cmd_sync_sheet(args):
    source = args.file or args.source
    if not source:
        sys.exit("Indiquer l'URL CSV publiée de l'onglet Décisions ou --file chemin.csv")
    rows = fetch_decisions(source)
    briefs = {b.id: b for b in load_briefs()}
    applied, ideas = 0, []
    for r in rows:
        bid = (r.get("id") or "").strip().upper()
        comment = (r.get("Commentaire") or "").strip()
        who = (r.get("Décidé par") or "").strip() or "feuille"
        when = (r.get("Date") or "").strip() or today()
        statut_raw = (r.get("Statut décidé") or "").strip()
        if not bid:
            if comment:
                ideas.append("%s (%s)" % (comment, who))
            continue
        if not ID_RE.match(bid):
            warn("id invalide ignoré : %s" % bid)
            continue
        b = briefs.get(bid)
        if b is None:
            warn("brief inconnu ignoré : %s" % bid)
            continue
        new = normalize_statut(statut_raw) if statut_raw else None
        change = new is not None and new != b.statut
        already = bool(comment) and comment in b.body
        if not change and (not comment or already):
            continue
        desc = "%s : %s%s" % (bid, ("statut %s → %s" % (b.statut, new)) if change else "commentaire",
                              (" — « %s »" % comment) if comment and not already else "")
        if args.dry_run:
            print("[simulation] " + desc)
            continue
        if change:
            old = b.statut
            b.data["statut"] = new
            b.append_journal("%s — %s (feuille) — statut : %s → %s" % (when, who, old, new))
        if comment and not already:
            b.append_journal("%s — %s (feuille) — %s" % (when, who, comment))
        b.path.write_text(dump_front(b.data, b.order) + b.body, encoding="utf-8")
        applied += 1
        print("Appliqué : " + desc)
    for w in warnings:
        print(w)
    if ideas:
        print("\nIdées à créer (lignes sans id) :")
        for i in ideas:
            print("  - " + i)
    print("\n%d décision(s) appliquée(s)." % applied)
    if applied:
        print("Relancer : python3 tools/plan.py build")


def main():
    ap = argparse.ArgumentParser(description="Plan éditorial CheckIA")
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("new")
    s.add_argument("slug")
    s.add_argument("--type", required=True)
    s.add_argument("--serie", default="")
    s.add_argument("--titre", required=True)
    s.add_argument("--mois", default="")
    s.add_argument("--auteur", default="")
    s.add_argument("--responsable", default="")
    s.set_defaults(func=cmd_new)
    s = sub.add_parser("build")
    s.add_argument("--no-xlsx", action="store_true")
    s.set_defaults(func=cmd_build)
    s = sub.add_parser("list")
    for opt in ("--statut", "--mois", "--serie", "--type", "--responsable"):
        s.add_argument(opt)
    s.set_defaults(func=cmd_list)
    s = sub.add_parser("show")
    s.add_argument("id")
    s.set_defaults(func=cmd_show)
    s = sub.add_parser("next")
    s.add_argument("--all", action="store_true")
    s.set_defaults(func=cmd_next)
    s = sub.add_parser("set")
    s.add_argument("id")
    s.add_argument("assignments", nargs="+")
    s.add_argument("--note")
    s.add_argument("--humain")
    s.add_argument("--build", action="store_true")
    s.set_defaults(func=cmd_set)
    s = sub.add_parser("check")
    s.set_defaults(func=cmd_check)
    s = sub.add_parser("sync-sheet")
    s.add_argument("source", nargs="?")
    s.add_argument("--file")
    s.add_argument("--dry-run", action="store_true")
    s.set_defaults(func=cmd_sync_sheet)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
